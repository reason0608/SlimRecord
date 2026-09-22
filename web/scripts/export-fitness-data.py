"""Export normalized workbook sheets to the static JSON used by the dashboard."""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SHEETS = ("DailyLogs", "FoodLogs", "ExerciseLogs", "Goals")


def serialize_value(value):
    """Convert Excel-compatible values to JSON-compatible values."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return value


def read_records(sheet):
    """Read a worksheet as records, skipping completely empty rows."""
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value) if value is not None else "" for value in rows[0]]
    records = []
    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue
        records.append(
            {
                header: serialize_value(value)
                for header, value in zip(headers, row)
                if header
            }
        )
    return records


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    missing = [name for name in SHEETS if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Missing worksheets: {', '.join(missing)}")

    data = {
        "dailyLogs": read_records(workbook["DailyLogs"]),
        "foodLogs": read_records(workbook["FoodLogs"]),
        "exerciseLogs": read_records(workbook["ExerciseLogs"]),
        "goals": read_records(workbook["Goals"]),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
